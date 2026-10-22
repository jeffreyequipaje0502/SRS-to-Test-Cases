"""
Reusable Excel generator for the SRS-to-TestCases skill.

This module exposes a single function, generate_test_cases_excel(), that
takes a list of test-case dictionaries and writes an Azure DevOps-importable
.xlsx file.  It is imported by per-SRS data scripts under output/.

Do NOT add test-case data to this file.  Each SRS gets its own data script:
    output/<SRS_name>_data.py

See output/SwiftCharge_example_data.py for a working example.
"""

import os

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def generate_test_cases_excel(test_cases, output_path, area_path=""):
    """Write test cases to an Azure DevOps-importable Excel file.

    Parameters
    ----------
    test_cases : list[dict]
        Each dict must have:
            - "title"       (str)   — test case title
            - "complexity"  (str)   — "1" (Simple), "3" (Medium), "5" (Complex)
            - "steps"       (list)  — list of {"action": ..., "expected": ...}
    output_path : str
        Destination .xlsx path (relative or absolute).
    area_path : str
        Azure DevOps area path, e.g. "PHL-CTB-CIB-08-IB".
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = [
        "ID", "Work Item Type", "Title", "Test Step",
        "Step Action", "Step Expected", "Area Path",
        "State", "Complexity", "Assigned To", "Status of automation"
    ]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 2
    for tc in test_cases:
        # Parent row
        ws.cell(row=row_num, column=1, value=None)
        ws.cell(row=row_num, column=2, value="Test Case")
        ws.cell(row=row_num, column=3, value=tc["title"])
        ws.cell(row=row_num, column=4, value=None)
        ws.cell(row=row_num, column=5, value=None)
        ws.cell(row=row_num, column=6, value=None)
        ws.cell(row=row_num, column=7, value=area_path)
        ws.cell(row=row_num, column=8, value="Ready")
        ws.cell(row=row_num, column=9, value=tc.get("complexity", "3"))
        ws.cell(row=row_num, column=10, value=None)
        ws.cell(row=row_num, column=11, value="Not yet automated")
        row_num += 1

        # Step rows
        for step_num, step in enumerate(tc["steps"], 1):
            ws.cell(row=row_num, column=4, value=step_num)
            ws.cell(row=row_num, column=5, value=step["action"])
            ws.cell(row=row_num, column=6, value=step["expected"])
            row_num += 1

    # Auto-size columns (capped at 80 chars for readability)
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        adjusted = min(max_len + 4, 80)
        ws.column_dimensions[col_letter].width = adjusted

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Saved: {output_path}")
